"""Prepare manuscript evidence and display items from supplied artifacts.

This script reads existing results. It does not train or evaluate a detector.
"""
from __future__ import annotations

import csv
import hashlib
import json
import pickletools
import shutil
import zipfile
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from loss_retrain_assets import prepare_pair, render_pair

PAPER = Path(__file__).resolve().parents[1]
ROOT = PAPER.parent
ARCHIVE = ROOT / 'Result/artifacts/ablation_analysis_20261012'
DOWNLOADS = Path('/home/duyennh/Downloads')
FIGURES = PAPER / 'figures'
RESEARCH = PAPER / 'research'
TABLES = PAPER / 'tables'
for directory in (FIGURES, RESEARCH, TABLES, PAPER / 'sections', RESEARCH / 'configs', FIGURES / 'pptx_media'):
    directory.mkdir(parents=True, exist_ok=True)

manifest = []
configuration_audit = []
def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def copy(source, target, purpose):
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, target)
    manifest.append({'source': str(source), 'target': str(target.relative_to(PAPER)),
                     'sha256': sha(source), 'purpose': purpose})

def resolve_run(row):
    original = Path(row['directory'])
    if (original / 'config.resolved.json').exists():
        return original
    for pattern in (row['name'], '*/' + row['name'], '*/*/' + row['name']):
        for candidate in DOWNLOADS.glob(pattern):
            if (candidate / 'config.resolved.json').exists():
                return candidate
    raise FileNotFoundError(row['name'])

inventory = json.loads((ARCHIVE / 'run_inventory.json').read_text())
labels = {'R01': 'B0', 'R02': 'O32', 'R03': 'H32', 'R04': 'RC24',
          'R05': 'O24G16', 'R06': 'O24G32', 'R07': 'S32', 'R08': 'O32G16',
          'R09': 'E14', 'R10': 'G16', 'R11': 'B1', 'R12': 'E8'}
runs = {}
summaries = []
for row in inventory:
    label = labels.get(row['run_id'], row['run_id'].replace('P0', 'P'))
    runs[label] = row
    run = resolve_run(row)
    evaluation = run / Path(row['evaluation']).name
    if not evaluation.exists():
        raise FileNotFoundError(evaluation)
    raw = json.loads(evaluation.read_text())
    recorded = raw['accuracy']['map_moderate_percent']
    if abs(recorded - row['map_moderate']) > 1e-8:
        raise ValueError(f'AP mismatch for {label}')
    if raw['model']['parameters'] != row['params']:
        raise ValueError(f'Parameter mismatch for {label}')
    provided_config = run / 'config.resolved.json'
    config = json.loads(provided_config.read_text())
    copy(provided_config, RESEARCH / 'configs' / 'as_provided' / (label + '.json'), 'Unmodified supplied config, including any post-run edits')
    if (run / 'run.json').exists():
        copy(run / 'run.json', RESEARCH / 'configs' / (label + '_run.json'), 'Actual run settings')
        recorded_hash = json.loads((run / 'run.json').read_text()).get('config_sha256')
    else:
        recorded_hash = None
    restored = False
    if recorded_hash and sha(provided_config) != recorded_hash:
        candidate = json.loads(provided_config.read_text())
        candidate['model'].pop('header_architecture', None)
        encoded = json.dumps(candidate, indent=2, sort_keys=True) + '\n'
        if hashlib.sha256(encoded.encode()).hexdigest() != recorded_hash:
            raise ValueError(f'Unresolved training config hash mismatch for {label}')
        config = candidate
        restored = True
    target = RESEARCH / 'configs' / (label + '.json')
    target.write_text(json.dumps(config, indent=2, sort_keys=True) + '\n')
    checkpoint = run / 'selected' / Path(row['checkpoint']).name
    head_keys = []
    if checkpoint.exists():
        # Inspect serialized parameter names without executing pickle instructions.
        with zipfile.ZipFile(checkpoint) as archive:
            data = next(n for n in archive.namelist() if n.endswith('data.pkl'))
            strings = [arg for op, arg, _ in pickletools.genops(archive.read(data))
                       if op.name in ('BINUNICODE', 'SHORT_BINUNICODE', 'UNICODE')]
            head_keys = [s for s in strings if 'header' in s and s.endswith('weight')][:12]
    configuration_audit.append({'label': label, 'supplied_sha256': sha(provided_config),
                                'recorded_training_sha256': recorded_hash,
                                'recovered_sha256': sha(target),
                                'removed_post_run_header_architecture': restored,
                                'checkpoint_head_parameter_names': head_keys})
    row['local_directory'] = str(run)
    row['paper_label'] = label
    for path in sorted(run.glob('evaluation*.json')):
        raw = json.loads(path.read_text())
        data = {k: v for k, v in raw.get('data', {}).items() if not isinstance(v, (dict, list))}
        summaries.append({'paper_label': label, 'source': str(path), 'sha256': sha(path),
                          'accuracy': raw['accuracy'], 'model': raw['model'],
                          'protocol': raw.get('protocol'), 'data': data,
                          'latency': raw.get('latency'), 'counts': raw.get('counts'),
                          'runtime': raw.get('runtime')})

final_run = Path((RESEARCH / 'final_3d_run_location.txt').read_text().strip())
final = {}
for mode in ('3d', 'bev'):
    source = final_run / f'evaluation_ap_{mode}.json'
    raw = json.loads(source.read_text())
    final[mode] = raw
    copy(source, RESEARCH / f'final_evaluation_ap_{mode}.json', 'Saved CUDA reference evaluation, 100 epoch 3D training')
copy(final_run / 'evaluation_loss_3d.json', RESEARCH / 'final_evaluation_loss_3d.json', 'Loss-selected checkpoint sensitivity')
copy(final_run / 'config.resolved.json', RESEARCH / 'configs' / 'Final3D.json', 'Trained shared head with vertical branch')
copy(final_run / 'run.json', RESEARCH / 'configs' / 'Final3D_run.json', 'Actual final training budget')
copy(final_run / 'selected/selection.json', RESEARCH / 'final_selection.json', 'Checkpoint selection on 3D R40 Moderate')
for name in ('summary.json', 'audit.json', 'failure_analysis.json', 'checkpoint_comparison.json', 'matched_objects.csv'):
    copy(ROOT / 'artifacts/simam_3d_diagnostics' / name, RESEARCH / ('diagnostic_' + name), 'Existing CPU diagnostic, distinct from reference AP')
for name in ('runs_summary.csv', 'operating_metrics.csv', 'operating_metrics_per_class.csv', 'pairwise_effects.csv',
             'paper_results_all_runs.csv', 'evidence_manifest.json'):
    copy(ARCHIVE / name, RESEARCH / name, 'Supplied ablation evidence')
for name in ('train', 'val'):
    copy(ROOT / f'Code/splits/kitti/{name}.txt', RESEARCH / f'{name}_ids.txt', 'Frame split used by manuscript')
train = set((RESEARCH / 'train_ids.txt').read_text().split())
val = set((RESEARCH / 'val_ids.txt').read_text().split())
if len(train) != 3712 or len(val) != 3769 or train & val:
    raise ValueError('Unexpected split or frame overlap')

# Keep the 50-epoch factorial intact and append the new matched loss pair.
loss_pair = prepare_pair(PAPER, copy)
runs.update(loss_pair['runs'])
inventory.extend(loss_pair['runs'].values())
summaries.extend(loss_pair['evaluations'])
configuration_audit.extend(loss_pair['configuration_audit'])

# Original PowerPoint remains intact; extract all raster resources and recursively inspect labels.
pptx = ROOT / 'Code/diagrams/LiDAR-MobilePixorNeXt_Diagrams.pptx'
copy(pptx, FIGURES / 'source_diagrams.pptx', 'User-supplied editable diagram source')
presentation = Presentation(pptx)
slides = []
def inspect(shapes, number, texts, pictures):
    for shape in shapes:
        if shape.shape_type == 6:
            inspect(shape.shapes, number, texts, pictures)
        if getattr(shape, 'has_text_frame', False) and shape.text.strip():
            texts.append(shape.text)
        if shape.shape_type == 13:
            name = f'slide_{number:02d}_{len(pictures)+1:02d}.{shape.image.ext}'
            (FIGURES / 'pptx_media' / name).write_bytes(shape.image.blob)
            pictures.append({'file': name, 'size': shape.image.size})
for number, slide in enumerate(presentation.slides, 1):
    texts, pictures = [], []
    inspect(slide.shapes, number, texts, pictures)
    slides.append({'slide': number, 'texts': texts, 'pictures': pictures})
(RESEARCH / 'pptx_inventory.json').write_text(json.dumps(slides, ensure_ascii=False, indent=2))

# Supplied vector counterparts preserve text and provide configuration-consistent labels.
for source, name, purpose in [
    ('Code/diagrams/figure1_flow/overview_flow_no_heatmap.pdf', 'overview.pdf', 'Existing overall pipeline diagram'),
    ('Code/diagrams/mobilepixornext/hist14_focal_simam_f32_s42/architecture_bev_backbone_heads.pdf', 'architecture_bev.pdf', 'Vector counterpart of the SimAM architecture in PowerPoint; 32 channel fusion'),
    ('Code/diagrams/mobilepixornext/head_current_vs_proposed/comparison.pdf', 'head_comparison.pdf', 'Supplied independent versus shared head schematic'),
    ('Code/diagrams/mobilepixornext/loss_attention/oga.pdf', 'oga.pdf', 'Existing oriented geometry objective schematic'),
    ('Code/diagrams/mobilepixornext/loss_attention/focal_context.pdf', 'focal_context.pdf', 'Existing Focal Context schematic'),
    ('Demo/outputs/bev/004338_validation/overview_preview.png', 'validation_bev.png', 'Validation frame 004338, P4 BEV predictions'),
    ('Demo/outputs/3d_lidar_lift/004338_validation/overview_preview.png', 'validation_lift.png', 'Validation BEV footprints with LiDAR-estimated vertical extent'),
    ('Demo/outputs/pipeline/005663/pipeline_overview.png', 'training_pipeline.png', 'Augmented training-frame illustration, P4 checkpoint'),
    ('Demo/outputs/pipeline/005663/04_hist14_channels.png', 'hist14_channels.png', 'All supplied Hist14 channels'),
    ('Demo/outputs/video/camera_lidar_bev_heatmap/contact_sheet.png', 'sequence.png', '174-frame validation video examples; RGB display only'),
    ('Demo/outputs/heatmap/005663/heatmap_preview.png', 'heatmap.png', 'Predicted classification probabilities and Gaussian training targets'),
]:
    copy(ROOT / source, FIGURES / name, purpose)

def ap_row(label):
    r = runs[label]
    return [f'{r["class_ap"][c]["Moderate"]:.2f}' for c in ('Car', 'Pedestrian', 'Cyclist')] + [f'{r["map_moderate"]:.2f}']

def table(name, columns, headers, rows):
    body = ['\\begin{tabular}{' + columns + '}', '\\toprule', ' & '.join(headers) + r' \\', '\\midrule']
    body.extend(' & '.join(str(x) for x in row) + r' \\' for row in rows)
    body.extend(['\\bottomrule', '\\end{tabular}'])
    (TABLES / name).write_text('\n'.join(body) + '\n')

table('loss_factorial.tex', 'llrrrr', ['Run', 'Objective / quality', 'Car', 'Ped.', 'Cyc.', 'mAP'],
      [[label, desc] + ap_row(label) for label, desc in [
          ('P1', 'Baseline'), ('P2', 'OGA'), ('P3', 'Baseline + IQA'), ('P4', 'OGA + IQA')]])
table('attention.tex', 'llrrr', ['Run', 'Local module', 'Ped. AP', 'mAP', 'Model ms'],
      [[label, desc, ap_row(label)[1], ap_row(label)[3], f'{runs[label]["latency_ms"]["model"]:.3f}']
       for label, desc in [('P4', 'None'), ('P14', 'Detail path'), ('P15', 'ECA'), ('P16', 'SimAM')]])
table('final_reference.tex', 'lrrrrrrr', ['Class', 'GT Mod.', 'BEV E', 'BEV M', 'BEV H', '3D E', '3D M', '3D H'],
      [[c, final['3d']['accuracy']['R40']['valid_gt_counts'][c]['Moderate']] +
       [f'{final[m]["accuracy"]["R40"]["per_class"][c][d]:.2f}' for m in ('bev','3d') for d in ('Easy','Moderate','Hard')]
       for c in ('Car', 'Pedestrian', 'Cyclist')] +
      [['Macro mean', '--'] + [f'{np.mean([final[m]["accuracy"]["R40"]["per_class"][c][d] for c in ("Car","Pedestrian","Cyclist")]):.2f}'
                               for m in ('bev','3d') for d in ('Easy','Moderate','Hard')]])
with (ARCHIVE / 'operating_metrics.csv').open(encoding='utf-8-sig') as stream:
    operating = list(csv.DictReader(stream))
(RESEARCH / 'operating_column_names.json').write_text(json.dumps(list(operating[0]), indent=2))
table('all_ap.tex', 'lrrrrrrrrrr', ['Run', 'Car E', 'Car M', 'Car H', 'Ped. E', 'Ped. M', 'Ped. H', 'Cyc. E', 'Cyc. M', 'Cyc. H', 'mAP M'],
      [[label] + [f'{r["class_ap"][c][d]:.2f}' for c in ('Car','Pedestrian','Cyclist') for d in ('Easy','Moderate','Hard')] + [f'{r["map_moderate"]:.2f}']
       for label, r in runs.items()])
table('all_runtime.tex', 'lrrrrrr', ['Run', 'Epochs', 'Selected', 'Parameters', 'Model ms', 'Pipeline ms', 'Pipeline Hz'],
      [[label, r['epochs_completed'], r['selected_epoch'], f'{r["params"]:,}', f'{r["latency_ms"]["model"]:.3f}',
        f'{r["latency_ms"]["input_to_detections"]:.3f}', f'{r["end_to_end_hz"]:.2f}'] for label, r in runs.items()])
table('checkpoint.tex', 'lrrrrr', ['Run', 'AP ep.', 'AP score', 'Loss ep.', 'Loss score', '$\\Delta$ pp'],
      [[label, runs[label]['selected_epoch'], f'{runs[label]["map_moderate"]:.2f}', runs[label]['loss_checkpoint_epoch'],
        f'{runs[label]["loss_checkpoint_map"]:.2f}', f'{runs[label]["map_moderate"]-runs[label]["loss_checkpoint_map"]:+.4f}']
       for label in ('S32','P4','P11','P16','RC24','O100','Q100')])

plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 8, 'axes.spines.top': False,
                     'axes.spines.right': False, 'pdf.fonttype': 42, 'svg.fonttype': 'none'})
def save(fig, name):
    fig.savefig(FIGURES / (name + '.pdf'), bbox_inches='tight')
    fig.savefig(FIGURES / (name + '.png'), dpi=300, bbox_inches='tight')
    fig.savefig(FIGURES / (name + '.svg'), bbox_inches='tight')
    plt.close(fig)

render_pair(loss_pair, table, save, PAPER, runs)

fig, axes = plt.subplots(1, 2, figsize=(7.0, 2.6), layout='constrained')
for quality, color, marker in [(False, '#0072B2', 'o'), (True, '#D55E00', 's')]:
    labels2 = ['P1','P2'] if not quality else ['P3','P4']
    values = [runs[l]['map_moderate'] for l in labels2]
    axes[0].plot([0,1], values, color=color, marker=marker, label='IQA on' if quality else 'IQA off')
    for x,y in enumerate(values): axes[0].annotate(f'{y:.2f}', (x,y), xytext=(0,6 if quality else -12), textcoords='offset points', ha='center')
axes[0].set(xticks=[0,1], xticklabels=['Baseline loss','OGA'], ylim=(66.7,71.4), ylabel='Local BEV mAP Moderate (%)', title='(a) OGA and quality interaction')
axes[0].legend(frameon=False, loc='center right', fontsize=7)
axes[0].grid(axis='y', alpha=.2)
for label in ('S32','P1','P2','P3','P4','P14','P15','P16'):
    r=runs[label]; x=r['latency_ms']['model'];y=r['map_moderate']
    color='#D55E00' if label=='P16' else '#0072B2'
    axes[1].scatter(x,y,s=25,color=color)
    offsets={'P3':(-22,5),'P4':(-18,5),'P14':(-20,-11),'P15':(5,-6),'P16':(5,0),'P1':(5,-5),'P2':(5,-2),'P11':(-18,7),'S32':(5,3)}
    axes[1].annotate(label,(x,y),xytext=offsets[label],textcoords='offset points',fontsize=7)
axes[1].set(xlabel='Model forward latency (ms)', ylabel='Local BEV mAP Moderate (%)',title='(b) Saved L4 / FP32 measurements',xlim=(6.75,8.9),ylim=(66.7,71.4))
axes[1].grid(alpha=.2)
save(fig,'ablation_tradeoff')

fig, axes=plt.subplots(1,2,figsize=(7.0,2.5),layout='constrained')
stages=('preprocess','host_to_device','model','decode_nms')
left=np.zeros(3)
for stage,color,name in zip(stages,['#56B4E9','#009E73','#0072B2','#E69F00'],['Load / encode','Transfer','Model','Decode / NMS']):
    values=np.array([runs[l]['latency_ms'][stage] for l in ('P3','P4','P16')]);axes[0].barh(['P3','P4','P16'],values,left=left,color=color,label=name);left+=values
axes[0].set(xlabel='Stage latency (ms)',title='(a) Offline pipeline components')
axes[0].legend(frameon=False,fontsize=6,ncol=2,loc='lower right')
positions=np.arange(3)
for i,(label,color) in enumerate([('P4','#0072B2'),('P16','#D55E00')]):
    values=[runs[label]['distance_ap']['Pedestrian'][band]['ap_r40_percent'] for band in ('0_30m','30_50m','50_70.4m')]
    axes[1].bar(positions+(i-.5)*.32,values,width=.32,label=label,color=color)
axes[1].set(xticks=positions,xticklabels=['0–30 m\n(n=1589)','30–50 m\n(n=137)','50–70.4 m\n(n=13)'],ylabel='Pedestrian AP Moderate (%)',title='(b) Range-stratified local BEV AP')
axes[1].legend(frameon=False,fontsize=7);axes[1].grid(axis='y',alpha=.2)
save(fig,'latency_range')

diag=json.loads((RESEARCH / 'diagnostic_summary.json').read_text())
audit=json.loads((RESEARCH / 'diagnostic_audit.json').read_text())
fig,axes=plt.subplots(1,2,figsize=(7.0,2.6),layout='constrained')
curve=audit['val_ap/score_r40_moderate']
axes[0].plot([x['step'] for x in curve],[x['value'] for x in curve],marker='o',color='#0072B2')
axes[0].axvline(70,color='#D55E00',ls='--',lw=.8);axes[0].set(xlabel='Training epoch',ylabel='Reference 3D mAP Moderate (%)',title='(a) Shared head 3D training')
axes[0].grid(alpha=.2)
classes=('Car','Pedestrian','Cyclist'); pos=np.arange(3)
for i,(key,color,display) in enumerate([('iou3d_pass_count','#0072B2','Predicted z and h'),('oracle_z_pass_count','#009E73','Oracle vertical location'),('oracle_zh_pass_count','#E69F00','Oracle location + height')]):
    vals=[100*diag['vertical_errors'][c][key]/diag['vertical_errors'][c]['bev_matched_objects'] for c in classes]
    axes[1].bar(pos+(i-1)*.24,vals,width=.24,color=color,label=display)
axes[1].set(xticks=pos,xticklabels=classes,ylim=(0,111),ylabel='3D IoU passes / fixed BEV matches (%)',title='(b) Conditional geometric diagnostic')
axes[1].legend(frameon=False,fontsize=6,loc='lower left');axes[1].grid(axis='y',alpha=.2)
save(fig,'vertical_diagnostic')

(RESEARCH / 'evaluation_summaries.json').write_text(json.dumps(summaries,ensure_ascii=False,indent=2))
(RESEARCH / 'run_inventory_resolved.json').write_text(json.dumps(inventory,ensure_ascii=False,indent=2))
(RESEARCH / 'asset_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
(RESEARCH / 'configuration_audit.json').write_text(json.dumps(configuration_audit, indent=2))
(RESEARCH / 'evidence_checks.json').write_text(json.dumps({
    'raw_ablation_runs_checked':len(inventory),'saved_evaluations_read':len(summaries),
    'original_archived_runs':len(inventory)-len(loss_pair['runs']),
    'original_archived_evaluations':len(summaries)-len(loss_pair['evaluations']),
    'matched_loss_retraining_runs':len(loss_pair['runs']),
    'matched_loss_retraining_evaluations':len(loss_pair['evaluations']),
    'matched_loss_configs_differ_only_in_objective_name':True,
    'matched_loss_retraining_delta_mAP_pp':loss_pair['comparison']['delta_Q_minus_O_pp']['map_moderate'],
    'completed_50_epoch_AP_selected_runs':sum(r['epochs_completed']==50 and r['selection']=='ap' for r in inventory),
    'train_frames':len(train),'validation_frames':len(val),'frame_overlap':len(train&val),
    'final_3d_parameters':final['3d']['model']['parameters'],
    'final_checkpoint_sha256':final['3d']['model']['sha256'],
    'final_3d_map_R40':final['3d']['accuracy']['R40']['map_moderate_percent'],
    'final_bev_map_R40':final['bev']['accuracy']['R40']['map_moderate_percent'],
    'new_training_or_evaluation_performed':False},indent=2))
print('Prepared 21 archived runs plus 2 matched loss retraining runs, 45 saved BEV reports, reference results, and manuscript display items.')
